# -*- coding: utf-8 -*-
"""由 evalset/items.py 生成 eval_set_50.jsonl / .md / .xlsx
- 根据 evidence 在 chunks（output/chunks_heading.jsonl + uploads/chunks_uploaded.jsonl）中定位 chunk_id、section(heading_path)
- 在 PDF 原件该 chunk 页码范围内搜索 evidence，得到精确页码 page（找不到则用 chunk 页码范围）
匹配规则：去掉全部空白后做子串匹配（PDF 抽取的换行/空格位置不稳定，不影响字符内容）"""
from __future__ import annotations
import glob, json, os, re, sys
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "evalset"))
from items import ITEMS, REFUSALS

def squash(s: str) -> str:
    return re.sub(r"\s+", "", s)

def load_chunks():
    cs = []
    for f in ["output/chunks_heading.jsonl", "uploads/chunks_uploaded.jsonl"]:
        p = os.path.join(HERE, f)
        if os.path.exists(p):
            cs += [json.loads(l) for l in open(p, encoding="utf-8") if l.strip()]
    return [c for c in cs if c["file_type"] == "pdf"]

def pdf_path(source):
    for p in glob.glob("/workspace/power_docs/**/" + source, recursive=True) + glob.glob(os.path.join(HERE, "uploads/files", source)):
        return p

def exact_page(source, ev, p0, p1):
    import pymupdf as fitz
    p = pdf_path(source)
    if not p:
        return None
    import unicodedata
    nz = lambda t: squash(unicodedata.normalize("NFKC", t))
    d, e = fitz.open(p), nz(ev)
    texts = {i: nz(d[i - 1].get_text()) for i in range(max(1, p0), min(len(d), p1) + 1)}
    for n in (14, 10, 7):                                    # 先长后短，取首个命中页
        for i, t in texts.items():
            if e[:n] in t:
                return i
    return None

def resolve(it, chunks):
    srcs = sorted({c["source"] for c in chunks if c["source"].startswith(it["source"])})
    assert len(srcs) == 1, (it["id"], it["source"], srcs)
    src, ev = srcs[0], squash(it["evidence"])
    hits = [c for c in chunks if c["source"] == src and ev in squash(c["text"])]
    assert hits, f"{it['id']}: evidence 不在 {src} 的任何 chunk 中"
    c = hits[0]
    pg = exact_page(src, it["evidence"], c["page_start"], c["page_end"])
    page = str(pg) if pg else (f"{c['page_start']}" + (f"-{c['page_end']}" if c["page_end"] != c["page_start"] else ""))
    if it.get("page_override"):                               # 扫描件（无文本层，OCR 文本见 output/parsed）人工核对页码
        page = it["page_override"]
    return {**{k: v for k, v in it.items() if k not in ("source", "page_override")}, "lang": it.get("lang", "zh"), "source": src,
            "section": it.get("section") or c["heading_path"], "chunk_heading": c["heading_path"], "page": page, "chunk_id": c["chunk_id"], "all_chunk_ids": [h["chunk_id"] for h in hits]}

FIELDS = ["id", "category", "difficulty", "lang", "device", "question", "reference_answer", "key_points", "source", "section", "page", "evidence", "chunk_id", "chunk_heading", "key_points_zh"]

def main():
    chunks = load_chunks()
    rows = [resolve(it, chunks) for it in ITEMS]
    rows += [{**r, "lang": "en" if re.match(r"[A-Za-z]", r["question"]) else "zh", "source": "", "section": "", "page": "", "evidence": "", "chunk_id": ""} for r in REFUSALS]
    with open(os.path.join(HERE, "eval_set_50.jsonl"), "w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps({k: r.get(k) for k in FIELDS + ["all_chunk_ids"] if k in r}, ensure_ascii=False) + "\n")
    # markdown
    L = ["# 运维问答评测集（50 题 + 5 道库外拒答）", "",
         "由 `build_eval_set.py` 从 `evalset/items.py` 生成；evidence 均为 chunk 原文子串（去空白比较），由 `validate_eval_set.py` 校验。", ""]
    for cat in ["查参数", "故障排查", "操作步骤", "拒答"]:
        sub = [r for r in rows if r["category"] == cat]
        L += [f"## {cat}（{len(sub)} 题）", ""]
        for r in sub:
            L += [f"### {r['id']}｜{r['difficulty']}｜{r['device']}{'｜EN' if r['lang'] == 'en' else ''}", "",
                  f"**问题**：{r['question']}", "", f"**标准答案**：{r['reference_answer']}", "",
                  "**评分要点**：" + "；".join(r["key_points"]), ""] + ([("**评分要点（中文，用于自动评分）**：" + "；".join(r["key_points_zh"])), ""] if r.get("key_points_zh") else [])
            if r["source"]:
                L += [f"**出处**：{r['source']}｜{r['section']}｜p{r['page']}", "", f"> 原文：{r['evidence']}", ""]
    open(os.path.join(HERE, "eval_set_50.md"), "w", encoding="utf-8").write("\n".join(L))
    # xlsx：按类别分 sheet + 汇总
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook(); ws0 = wb.active; ws0.title = "汇总"
    ws0.append(["类别", "题数", "英文题"])
    for cat in ["查参数", "故障排查", "操作步骤", "拒答"]:
        sub = [r for r in rows if r["category"] == cat]
        ws0.append([cat, len(sub), sum(r["lang"] == "en" for r in sub)])
        ws = wb.create_sheet(cat)
        cols = ["id", "difficulty", "lang", "device", "question", "reference_answer", "key_points", "source", "section", "page", "evidence", "chunk_id"]
        ws.append(cols)
        for r in sub:
            ws.append([("\n".join(r[c]) if isinstance(r[c], list) else r[c]) for c in cols])
        for c, w in zip("ABCDEFGHIJKL", [6, 6, 5, 12, 40, 70, 40, 30, 30, 6, 40, 30]):
            ws.column_dimensions[c].width = w
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
    wb.save(os.path.join(HERE, "eval_set_50.xlsx"))
    print(f"写出 {len(rows)} 题")

if __name__ == "__main__":
    main()
