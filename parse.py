# -*- coding: utf-8 -*-
"""统一文档解析：PDF（PyMuPDF）与 DOCX（python-docx）-> ParsedDoc

用法：
    from parse import parse_document
    doc = parse_document("xx.pdf", max_pages=None)
    doc.pages -> [{"page": 1, "text": "...", "chars": 123, "scanned": False}]
命令行：
    python parse.py <file> [--max-pages N] [--out parsed.json]
"""
from __future__ import annotations
import argparse, json, os, re, unicodedata
from collections import Counter
from dataclasses import dataclass, field, asdict

import pymupdf  # PyMuPDF
import docx  # python-docx
from docx.oxml.ns import qn

# ---------------- 数据结构 ----------------
@dataclass
class ParsedDoc:
    source: str                      # 文件名
    path: str
    file_type: str                   # pdf / docx
    n_pages: int                     # 原文档页数
    pages: list = field(default_factory=list)   # [{"page", "text", "chars", "scanned"}]
    scanned_pages: list = field(default_factory=list)
    is_scanned: bool = False         # 大部分页无文本层 -> 需 OCR
    removed_header_footer: list = field(default_factory=list)  # 被剔除的页眉页脚模板（示例）
    docx_headings: list = field(default_factory=list)          # docx 中样式为 Heading 的段落
    font_headings: list = field(default_factory=list)          # PDF 中按字号/粗体识别的标题候选 [{page,text,size,level}]
    toc_pages: list = field(default_factory=list)              # 目录页
    ocr_pages: list = field(default_factory=list)              # 经 OCR 识别的页 [{page, avg_conf, rows, low_conf_boxes}]
    tables: list = field(default_factory=list)                 # 结构化表格（tables=True 时），见 tables.py

    @property
    def text(self) -> str:
        return "\n".join(p["text"] for p in self.pages)

    def to_dict(self):
        d = asdict(self)
        d["total_chars"] = sum(p["chars"] for p in self.pages)
        return d

# ---------------- 通用清洗 ----------------
CJK = r"\u3400-\u9fff\uf900-\ufaff"
SENT_END = "。！？；：:;!?」』）)】"
PAGE_NUM_RE = re.compile(
    r"^\s*(?:[-—–]\s*)?(?:第\s*)?(?:[0-9]{1,4}|[IVXLCivxlc]{1,6}|[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]{1,3})(?:\s*页)?"
    r"(?:\s*[/,]?\s*共\s*\d+\s*页|\s*/\s*\d+)?\s*(?:[-—–])?\s*$"
    r"|^\s*(?:Page|页码)\s*\d+(?:\s*(?:of|/)\s*\d+)?\s*$", re.I)
LIST_START_RE = re.compile(
    r"^\s*(?:[a-zA-Z][)）]|[（(]\d+[)）]|\d+[)）]|[①-⑳]|——|—|•|·|-\s|\*|[一二三四五六七八九十]+、|"
    r"\d+(?:\.\d+)*\s|[A-Z]\.\d+|附\s*录|表\s*\d|图\s*\d|Table\s|Fig)")
HEADING_LIKE_RE = re.compile(
    rf"^\s*(?:第[一二三四五六七八九十百\d]+[章节条]|\d{{1,2}}(?:\.\d{{1,2}}){{0,4}}\.?\s+\S|[A-Z]\.\d+(?:\.\d+)*\s+\S|附\s*录\s*[A-Z])")

def norm_line(s: str) -> str:
    s = unicodedata.normalize("NFKC", s)
    s = s.replace("\u00a0", " ").replace("\t", " ")
    s = re.sub(r"[\uf000-\uf8ff]", "", s)          # 私有区字符（公式符号字体）
    s = re.sub(r"[ \u3000]{2,}", " ", s)
    return s.strip()

def template_key(s: str) -> str:
    """页眉页脚模板：数字替换为 #，便于跨页比较"""
    return re.sub(r"\d+", "#", re.sub(r"\s+", "", s))[:80]

def is_cjk(ch: str) -> bool:
    return bool(ch) and re.match(f"[{CJK}]", ch) is not None


SOFT = "\x00"   # 版面块边界（软换行，可能是断行也可能是段落）
TOC_RE = re.compile(r"(\.{4,}|…{2,}|·{3,}|-{4,}|_{4,})\s*[\dIVX]*\s*$")
UNIT_RE = re.compile(r"^(kV|V|A|kA|mm|cm|m|s|ms|min|h|Hz|%|°C|℃|MPa|kPa|Pa|kg|t|mA|μA|Ω|MΩ|N|N·m|rad|mrad|K|mK|W|kW|MW|kVA)\b", re.I)
_CN_NUM = "一二三四五六七八九十百零"

def dedupe_repeated(s: str) -> str:
    """'ABCABCABC' -> 'ABC'（PDF 描边/阴影字造成的重复）"""
    m = re.fullmatch(r"(.{2,40}?)\1{1,4}", s)
    return m.group(1) if m else s

def detect_heading(line: str):
    """识别编号标题。返回 dict(number, title, level, kind) 或 None。
    支持：'1 范围' '5.2.1 xxx' '1范围' '第三章 xxx' '附录 A（规范性）xxx' 'A.1 xxx' '3.1 Title' '3.1. Title'"""
    s = line.strip()
    if not s or len(s) > 90 or TOC_RE.search(s):
        return None
    m = re.match(rf"^第\s*([{_CN_NUM}\d]+)\s*([章节条])\s*(.{{0,40}})$", s)
    if m and (not m.group(3).strip() or (_title_ok(m.group(3)) and not m.group(3).strip().startswith(("的", "中", "规定", "要求")))):
        lvl = {"章": 1, "节": 2, "条": 3}[m.group(2)]
        return {"number": f"第{m.group(1)}{m.group(2)}", "title": m.group(3).strip(), "level": lvl, "kind": "cn_chapter"}
    m = re.match(r"^附\s*录\s*([A-Z])\s*(?:[（(]\s*(?:规范性|资料性|normative|informative)\s*[)）])?\s*(.{0,40})$", s)
    if m and (not m.group(2).strip() or _title_ok(m.group(2))):
        return {"number": f"附录{m.group(1)}", "title": m.group(2).strip(), "level": 1, "kind": "appendix"}
    m = re.match(r"^([A-Z])\.(\d{1,2}(?:\.\d{1,2}){0,3})\s+(\S.{0,40})$", s)
    if m and _title_ok(m.group(3)):
        num = f"{m.group(1)}.{m.group(2)}"
        return {"number": num, "title": m.group(3).strip(), "level": num.count(".") + 1, "kind": "appendix_clause"}
    m = re.match(r"^(\d{1,2}(?:\.\d{1,2}){0,4})(\.?)(\s*)(\S.*)$", s)
    if m:
        num, dot, sp, title = m.groups()
        if not sp and not re.match(f"[{CJK}]", title):   # '1范围' 允许无空格；英文必须有空格
            return None
        if not _title_ok(title):
            return None
        latin = not re.search(f"[{CJK}]", title)
        if latin:
            if not re.match(r"[A-Z]", title) or title.endswith((".", ",", ";")) or len(title) > 80 or len(title.split()) > 10:
                return None
        level = num.count(".") + 1
        if level == 1 and int(num) > 30:
            return None
        return {"number": num, "title": title.strip(), "level": level, "kind": "numbered_en" if latin else "numbered_cn"}
    return None

def _title_ok(title: str) -> bool:
    t = title.strip()
    if not t or UNIT_RE.match(t) or len(t) < 2:
        return False
    if re.fullmatch(r"[A-Z0-9][A-Z0-9.\-_/]*\d[A-Z0-9.\-_/]*", t):   # 文档编号/图号，如 OPG.412.866-1、S1B7039801-03
        return False
    if re.search(r"[，,；;。！？]", t):        # 标题一般不含句内/句末标点
        return False
    if re.match(r"[^\w\u3400-\u9fff（(]", t):    # 以标点开头（如 '附录A;'）
        return False
    if re.search(r"[:：]", t) and len(t) > 30:     # 带冒号的长行多为条款正文
        return False
    cjk = len(re.findall(f"[{CJK}]", t))
    if cjk and cjk > 30:
        return False
    if re.fullmatch(r"[\d\s.\-—/×%~±+<>≤≥=]+", t):   # 纯数字/表格数值
        return False
    return True

def merge_broken_lines(lines: list[str]) -> list[str]:
    """断行合并：上一行未以句末标点结束、且不是标题；下一行不是列表/标题起始 -> 合并。
    中文之间直接拼接；英文之间加空格；英文连字符断词去掉连字符。
    SOFT（版面块边界）在满足合并条件时被吞掉，否则视为段落分隔。"""
    out: list[str] = []
    pending_soft = False
    for ln in lines:
        if ln == SOFT:
            pending_soft = True
            continue
        if not ln:
            if out and out[-1] != "":
                out.append("")
            pending_soft = False
            continue
        ln = dedupe_repeated(ln)
        if out and out[-1]:
            h = detect_heading(out[-1])
            if h and h["title"] == ln:            # 标题文字的阴影重复
                continue
            prev = out[-1]
            can_merge = (prev[-1] not in SENT_END and not detect_heading(prev)
                         and not LIST_START_RE.match(ln) and not detect_heading(ln)
                         and not PAGE_NUM_RE.match(ln) and len(prev) >= 15
                         and not TOC_RE.search(prev) and "⟦TABLE" not in prev and not ln.startswith("⟦TABLE")
                         and " | " not in prev and " | " not in ln)      # OCR 表格行不合并
            if can_merge:
                if prev.endswith("-") and re.match(r"[a-z]", ln):
                    out[-1] = prev[:-1] + ln
                elif is_cjk(prev[-1]) or is_cjk(ln[0]) or prev[-1] in "，、（(":
                    out[-1] = prev + ln
                else:
                    out[-1] = prev + " " + ln
                pending_soft = False
                continue
            if pending_soft:
                out.append("")
        out.append(ln)
        pending_soft = False
    res, blank = [], False
    for ln in out:
        if ln == "":
            if not blank and res:
                res.append("")
            blank = True
        else:
            res.append(ln); blank = False
    while res and res[-1] == "":
        res.pop()
    return res

PAGE_MARK_RE = re.compile(r"^\s*(?:第\s*\d+\s*页\s*)?(?:[,，/]?\s*共\s*\d+\s*页)?\s*$")

def is_page_mark(s: str) -> bool:
    """显式的 '第 2 页' / '共 11 页' / '第2页 共11页'，无论在页面什么位置都剔除"""
    return bool(s.strip()) and "页" in s and bool(PAGE_MARK_RE.match(s))

# ---------------- PDF ----------------
def _column_order(blocks, width):
    """双栏版面重排：整页宽的块作为分隔，分隔之间的块先左栏(自上而下)再右栏。
    仅当左右两栏各有 ≥3 个较长文本块时启用（避免打乱单栏页面和表格）。"""
    mid = width / 2
    col = lambda b: 0 if b[2] <= mid + 15 else (1 if b[0] >= mid - 15 else -1)
    left = sum(1 for b in blocks if col(b) == 0 and len(b[4]) > 60)
    right = sum(1 for b in blocks if col(b) == 1 and len(b[4]) > 60)
    if left < 3 or right < 3:
        return blocks
    out, band = [], []
    for b in sorted(blocks, key=lambda b: (b[1], b[0])):
        if col(b) == -1:
            out += sorted(band, key=lambda x: (col(x), x[1], x[0])); band = []
            out.append(b)
        else:
            band.append(b)
    return out + sorted(band, key=lambda x: (col(x), x[1], x[0]))

TABLE_PH = "⟦TABLE:{}⟧"
TABLE_PH_RE = re.compile(r"⟦TABLE:([^⟧]+)⟧")

def _in_rects(b, rects, ratio=0.5) -> bool:
    area = max(1e-6, (b[2] - b[0]) * (b[3] - b[1]))
    for r in rects:
        w = min(b[2], r[2]) - max(b[0], r[0]); h = min(b[3], r[3]) - max(b[1], r[1])
        if w > 0 and h > 0 and w * h >= ratio * area:
            return True
    return False

def _insert_placeholders(blocks, tables):
    """在阅读顺序中表格所在位置插入占位行（切块阶段据此确定表格所在章节，并从正文中移除）"""
    for t in sorted(tables, key=lambda t: t["bbox"][1]):
        ph = (t["bbox"][1], t["bbox"][3], [TABLE_PH.format(t["table_id"])])
        pos = next((k for k, b in enumerate(blocks) if b[0] >= t["bbox"][1]), len(blocks))
        blocks.insert(pos, ph)
    return blocks

def _pdf_blocks(page, exclude=None):
    """返回 [(y0, y1, [lines])]，按阅读顺序；exclude：表格区域（落在其中的文本块不进正文流）"""
    blocks = []
    raw = [b for b in page.get_text("blocks", sort=True) if b[6] == 0]   # 去掉图片块
    if exclude:
        raw = [b for b in raw if not _in_rects(b, exclude)]
    # 编号与标题分属左右两个文本块（如 SIPROTEC：'1.1' | 'General'）-> 合并为一行
    used = set()
    num_re = re.compile(r"\s*(\d{1,2}(?:\.\d{1,3}){0,4}|[A-Z](?:\.\d{1,3}){1,4})\.?\s*")
    for i, b in enumerate(raw):
        if i in used or not num_re.fullmatch(b[4]):
            continue
        for j, c in enumerate(raw):
            if j != i and j not in used and abs(c[1] - b[1]) < 4 and 0 <= c[0] - b[2] < 80 \
                    and len([x for x in c[4].split("\n") if x.strip()]) == 1 \
                    and detect_heading(b[4].strip().rstrip(".") + " " + c[4].strip()):
                raw[i] = (b[0], min(b[1], c[1]), c[2], max(b[3], c[3]), b[4].strip().rstrip(".") + " " + c[4].strip(), b[5], 0)
                used.add(j)
                break
    raw = [b for k, b in enumerate(raw) if k not in used]
    raw = _column_order(raw, page.rect.width)
    for b in raw:
        lines = [norm_line(x) for x in b[4].split("\n")]
        lines = [x for x in lines if x]
        # '5' + 'Installation' 分两行排版的章标题 -> '5 Installation'
        joined = []
        for x in (lines if len(lines) == 2 else []):
            if joined and re.fullmatch(r"\d{1,2}(?:\.\d{1,2}){0,3}\.?", joined[-1]) and detect_heading(joined[-1].rstrip(".") + " " + x):
                joined[-1] = joined[-1].rstrip(".") + " " + x
            else:
                joined.append(x)
        lines = joined or lines
        if lines:
            blocks.append((b[1], b[3], lines))
    return blocks

ADMONITIONS = {"DANGER", "WARNING", "CAUTION", "NOTICE", "NOTE", "IMPORTANT", "危险", "警告", "小心", "注意", "注", "提示", "说明"}

def _font_heading_candidates(pages_dicts, body_size):
    """字号明显大于正文（或同字号粗体短行）的行 -> 标题候选，按字号分级"""
    cands = []
    for page_no, h, dct in pages_dicts:
        for b in dct["blocks"]:
            if b.get("type") != 0:
                continue
            for l in b["lines"]:
                spans = [sp for sp in l["spans"] if sp["text"].strip()]
                if not spans:
                    continue
                text = norm_line("".join(sp["text"] for sp in spans))
                size = max(sp["size"] for sp in spans)
                bold = all((sp["flags"] & 16) or "Bold" in sp["font"] for sp in spans)
                y0, y1 = l["bbox"][1], l["bbox"][3]
                if y1 < h * 0.06 or y0 > h * 0.94:
                    continue
                if not (2 <= len(text) <= 60) or re.fullmatch(r"[\d\s.,:;/()%-]+", text) or TOC_RE.search(text):
                    continue
                if text.endswith(("。", ".", ",", "，", ";", "；")) and not re.match(r"^\d", text):
                    continue
                if text.upper() in ADMONITIONS:          # 安全警示标签不是章节
                    continue
                if size >= body_size * 1.18 or (bold and size >= body_size * 1.05 and len(text) <= 40):
                    cands.append({"page": page_no, "text": text, "size": round(size, 1), "bold": bold})
    sizes = sorted({c["size"] for c in cands}, reverse=True)
    rank = {sz: min(i + 1, 3) for i, sz in enumerate(sizes)}
    for c in cands:
        c["level"] = rank[c["size"]]
    return cands

def _needs_ocr(blocks, has_img: bool) -> bool:
    chars = sum(len(re.sub(r"\s", "", x)) for _, _, ls in blocks for x in ls)
    return chars == 0 or (chars < 30 and has_img)

def parse_pdf(path: str, max_pages: int | None = None, start_page: int = 1,
              ocr: str = "auto", ocr_dpi: int = 200, tables: bool = False, ocr_engine: str | None = None) -> ParsedDoc:
    """ocr: auto（无文本层的页自动 OCR）/ off（只标记不识别）/ force（所有页都 OCR）
    tables: True -> 结构化抽取表格（文字页 PyMuPDF find_tables；OCR 页用 PP-StructureV3 表格识别），表格区域文字移出正文流
    ocr_engine: paddle / rapid / None（ocr.default_engine()）"""
    d = pymupdf.open(path)
    first = max(1, start_page) - 1
    last = d.page_count if not max_pages else min(d.page_count, first + max_pages)
    doc = ParsedDoc(source=os.path.basename(path), path=path, file_type="pdf", n_pages=d.page_count)

    raw = []          # (page_no, height, blocks, has_images)
    ocr_font_heads = []
    dicts = []
    size_counter = Counter()
    for i in range(first, last):
        p = d[i]
        blocks, has_img = _pdf_blocks(p), bool(p.get_images())
        page_tables = []
        if tables and not _needs_ocr(blocks, has_img) and ocr != "force":
            from tables import extract_pdf_tables
            page_tables = extract_pdf_tables(p, doc.source, start_idx=0)
            if page_tables:
                blocks = _insert_placeholders(_pdf_blocks(p, exclude=[t["bbox"] for t in page_tables]), page_tables)
        if ocr == "force" or (ocr == "auto" and _needs_ocr(blocks, has_img)):
            try:
                from ocr import ocr_page, ocr_rows_as_blocks
                res = ocr_page(p, dpi=ocr_dpi, engine=ocr_engine)
                if tables and res.get("tables"):
                    from tables import tables_from_ocr
                    page_tables = tables_from_ocr(res, doc.source)
                elif not tables and res.get("tables"):
                    from ocr import _rows, _vertical_lines
                    res = dict(res, rows=_rows(res["boxes"], _vertical_lines(p)))      # 不做表格结构化：表格文字照常进正文行
                if res["rows"] or page_tables:
                    blocks = ocr_rows_as_blocks(res)
                    if page_tables:
                        blocks = _insert_placeholders(blocks, page_tables)
                    for r in res["rows"]:     # OCR 行高明显大于本页中位数 -> 字号标题候选（供 big_font 判断 / font 模式）
                        if r.get("big") and len(r["text"]) <= 30 and " | " not in r["text"]:
                            ocr_font_heads.append({"page": i + 1, "text": r["text"], "size": r["big"], "bold": False,
                                                   "level": 1 if r["big"] >= 1.8 else 2, "ocr": True})
                doc.ocr_pages.append({"page": i + 1, "avg_conf": res["avg_conf"], "rows": len(res["rows"]),
                                      "boxes": res["n_boxes"], "low_conf_boxes": res["low_conf_boxes"],
                                      "engine": res.get("engine", "rapid"), "tables": len(res.get("tables", []))})
            except ImportError:
                ocr = "off"          # 未安装 OCR 引擎：退化为只标记
        doc.tables += page_tables
        raw.append((i + 1, p.rect.height, blocks, has_img))
        dct = p.get_text("dict")
        dicts.append((i + 1, p.rect.height, dct))
        for b in dct["blocks"]:
            for l in b.get("lines", []):
                for sp in l["spans"]:
                    size_counter[round(sp["size"], 1)] += len(sp["text"].strip())
    body_size = size_counter.most_common(1)[0][0] if size_counter else 10

    dropped_ph = set()
    if doc.tables:
        from tables import merge_continuations, drop_template_rows
        ids = [t["table_id"] for t in doc.tables]
        doc.tables = merge_continuations(drop_template_rows(doc.tables), {pn: h for pn, h, _, _ in raw})
        dropped_ph = set(ids) - {t["table_id"] for t in doc.tables}      # 续表并入前表：去掉其占位行

    # 1) 统计页眉页脚候选：位于页面上/下 8% 区域的行，模板跨页重复
    edge_counter = Counter()
    any_counter = Counter()
    for _, h, blocks, _ in raw:
        seen_edge, seen_any = set(), set()
        for y0, y1, lines in blocks:
            for ln in lines:
                k = template_key(ln)
                if not k:
                    continue
                if y1 < h * 0.08 or y0 > h * 0.92:
                    seen_edge.add(k)
                if len(ln) <= 80:
                    seen_any.add(k)
        edge_counter.update(seen_edge); any_counter.update(seen_any)
    n = max(1, len(raw))
    # 页眉页脚：上下边缘区的行模板在 ≥30% 页面出现，或（短行）在 ≥3 页出现（如按章变化的书眉 "Generalities"）
    edge_tpl_strong = {k for k, c in edge_counter.items() if c >= max(2, 0.3 * n)}
    edge_tpl_weak = {k for k, c in edge_counter.items() if c >= 3 and len(k) <= 60} - edge_tpl_strong
    edge_tpl = edge_tpl_strong | edge_tpl_weak
    kept_once = set()   # 按章变化的书眉：第一次出现通常就是章标题本身，保留
    # 页面任意位置重复出现的短行（如每页的版权声明、文档号框），出现于 ≥50% 页面
    any_tpl = {k for k, c in any_counter.items() if c >= max(3, 0.5 * n) and len(k) >= 4}
    doc.removed_header_footer = sorted(edge_tpl | any_tpl)[:30]

    for page_no, h, blocks, has_img in raw:
        lines = []
        for y0, y1, blines in blocks:
            in_edge = y1 < h * 0.08 or y0 > h * 0.92
            for ln in blines:
                m = TABLE_PH_RE.fullmatch(ln)
                if m:
                    if m.group(1) not in dropped_ph:
                        lines += [SOFT, ln, SOFT]
                    continue
                k = template_key(ln)
                if k in any_tpl or is_page_mark(ln):
                    continue
                if lines and ln == lines[-1]:          # 连续重复行（描边/阴影字）
                    continue
                if in_edge and k in edge_tpl_weak and k not in kept_once and not PAGE_NUM_RE.match(ln):
                    kept_once.add(k)
                elif in_edge and (k in edge_tpl or PAGE_NUM_RE.match(ln)):
                    continue
                if PAGE_NUM_RE.match(ln) and (in_edge or len(ln) <= 3 and (y0 > h * 0.85 or y1 < h * 0.12)):
                    continue
                lines.append(ln)
            lines.append(SOFT)   # 块边界 -> 软换行
        merged = merge_broken_lines(lines)
        text = "\n".join(merged)
        chars = len(re.sub(r"\s", "", text))
        ocr_info = next((o for o in doc.ocr_pages if o["page"] == page_no), None)
        scanned = ocr_info is not None or (chars < 30 and has_img or chars == 0)   # scanned = 无文本层
        n_leader = len(re.findall(r"(\.{6,}|…{3,}|·{5,})\s*[\dIVX]+", text))
        if any(x.replace(" ", "").lower() in ("目次", "目录", "contents", "tableofcontents") for x in merged[:15]) \
                or n_leader >= 6:          # 目录续页：点引导线密集
            doc.toc_pages.append(page_no)
        if scanned:
            doc.scanned_pages.append(page_no)
        pg = {"page": page_no, "text": text, "chars": chars, "scanned": scanned}
        if ocr_info:
            pg.update({"ocr": True, "ocr_conf": ocr_info["avg_conf"]})
        doc.pages.append(pg)
    doc.is_scanned = len(doc.scanned_pages) >= 0.5 * n
    tpl = set(doc.removed_header_footer)
    doc.font_headings = [c for c in _font_heading_candidates(dicts, body_size)
                         if template_key(c["text"]) not in tpl and c["page"] not in doc.toc_pages] + ocr_font_heads
    return doc

# ---------------- DOCX ----------------
def _iter_body(document):
    """按文档顺序遍历段落与表格"""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)

def _para_page_breaks(p, use_rendered: bool) -> int:
    """段落中的分页数：显式分页符 w:br type=page、分节符 sectPr（pdf2docx 每页一节）；
    若文档没有上述标记，则退而使用 Word 保存时写入的 lastRenderedPageBreak。"""
    el = p._p
    n = sum(1 for br in el.iter(qn("w:br")) if br.get(qn("w:type")) == "page")
    ppr = el.find(qn("w:pPr"))
    sect = ppr.find(qn("w:sectPr")) if ppr is not None else None
    if sect is not None:
        typ = sect.find(qn("w:type"))
        # continuous 分节（pdf2docx 用于分栏）不换页；缺省/nextPage/oddPage/evenPage 换页
        if typ is None or typ.get(qn("w:val")) in (None, "nextPage", "oddPage", "evenPage"):
            n += 1
    if use_rendered:
        n += len(list(el.iter(qn("w:lastRenderedPageBreak"))))
    return n

def parse_docx(path: str) -> ParsedDoc:
    document = docx.Document(path)
    doc = ParsedDoc(source=os.path.basename(path), path=path, file_type="docx", n_pages=0)
    body = document.element.body
    has_explicit = any(True for _ in body.iter(qn("w:sectPr"))) and len(list(body.iter(qn("w:sectPr")))) > 1 \
        or any(br.get(qn("w:type")) == "page" for br in body.iter(qn("w:br")))
    page = 1
    buf: dict[int, list[str]] = {1: []}
    from docx.table import Table
    for item in _iter_body(document):
        if isinstance(item, Table):
            for row in item.rows:
                cells = []
                for c in row.cells:
                    t = norm_line(c.text.replace("\n", " "))
                    if t and (not cells or cells[-1] != t):   # 合并单元格去重
                        cells.append(t)
                if cells:
                    buf.setdefault(page, []).append(" | ".join(cells))
            buf.setdefault(page, []).append("")
            continue
        t = norm_line(item.text)
        style = (item.style.name if item.style is not None else "") or ""
        if t:
            if style.lower().startswith("heading") or style.startswith("标题"):
                doc.docx_headings.append({"page": page, "text": t, "style": style})
            buf.setdefault(page, []).append(t)
        else:
            buf.setdefault(page, []).append("")
        nb = _para_page_breaks(item, use_rendered=not has_explicit)
        if nb:
            page += nb
            buf.setdefault(page, [])
    pages_all = sorted(buf)
    # 页眉页脚：docx 的页眉页脚在 section.header 中，正文不包含；但 PDF 转换来的 docx 会把页眉作为正文
    cnt = Counter()
    for p in pages_all:
        cnt.update({template_key(x) for x in buf[p] if x and len(x) <= 80})
    tpl = {k for k, c in cnt.items() if len(pages_all) >= 3 and c >= max(3, 0.5 * len(pages_all)) and len(k) >= 4}
    doc.removed_header_footer = sorted(tpl)[:30]
    for p in pages_all:
        lines = [x for x in buf[p] if not (x and (template_key(x) in tpl or PAGE_NUM_RE.match(x)))]
        merged = merge_broken_lines(lines)
        text = "\n".join(merged)
        chars = len(re.sub(r"\s", "", text))
        if chars == 0 and p != pages_all[-1]:
            doc.scanned_pages.append(p)
        doc.pages.append({"page": p, "text": text, "chars": chars, "scanned": chars == 0})
    doc.n_pages = len(pages_all)
    doc.is_scanned = sum(x["chars"] for x in doc.pages) == 0
    return doc

# ---------------- 统一入口 ----------------
def parse_document(path: str, max_pages: int | None = None, start_page: int = 1, ocr: str = "auto",
                   tables: bool = False, ocr_engine: str | None = None) -> ParsedDoc:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return parse_pdf(path, max_pages=max_pages, start_page=start_page, ocr=ocr, tables=tables, ocr_engine=ocr_engine)
    if ext == ".docx":
        return parse_docx(path)
    raise ValueError(f"不支持的格式: {ext}")

def pdf_to_docx(pdf_path: str, docx_path: str, start: int = 0, end: int | None = None):
    """用 pdf2docx 把 PDF 转成 docx（用于测试 docx 解析路径）"""
    from pdf2docx import Converter
    cv = Converter(pdf_path)
    try:
        cv.convert(docx_path, start=start, end=end)
    finally:
        cv.close()

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--max-pages", type=int, default=None)
    ap.add_argument("--out", default=None)
    ap.add_argument("--ocr", choices=["auto", "off", "force"], default="auto", help="无文本层页面是否 OCR")
    a = ap.parse_args()
    d = parse_document(a.file, a.max_pages, ocr=a.ocr)
    info = {k: v for k, v in d.to_dict().items() if k != "pages"}
    print(json.dumps(info, ensure_ascii=False, indent=1)[:2000])
    print("---- 前 800 字 ----")
    print(d.text[:800])
    if a.out:
        json.dump(d.to_dict(), open(a.out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
